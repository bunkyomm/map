#!/usr/bin/env python3
"""stories.xlsx (シート1) から spots.json を生成する。

使い方:
    python3 tools/build_json.py

再実行しても同じ入力からは同じ出力になる（冪等）。
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "stories.xlsx"
SHEET_NAME = "シート1"
IMAGES_DIR = ROOT / "images"
AUDIO_DIR = ROOT / "audio"
OUTPUT_PATH = ROOT / "spots.json"

# 文京区周辺の緯度経度の目安（この範囲外は警告）
LAT_MIN, LAT_MAX = 35.69, 35.74
LNG_MIN, LNG_MAX = 139.7, 139.8

EXPECTED_HEADER = [
    "ID", "年度", "テーマ", "タイトル", "作者", "場所", "本文",
    "longitude", "latitude", "写真1", "写真2", "写真3", "音声", "備考",
]


def cell_str(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_year(value, row_no, warnings):
    s = cell_str(value)
    digits = re.sub(r"[^0-9]", "", s)
    if not digits:
        warnings.append(f"行{row_no}: 年度が空、または数値に変換できません（値: {value!r}）")
        return None
    return int(digits)


def list_dir_names(dir_path):
    """ディレクトリ内のファイル名を実バイト列として取得する。

    macOS(APFS)は既定で大文字小文字を区別しないため、Path.exists() では
    'a.JPG' を参照していても実体が 'a.jpg' であるケースを検出できない。
    GitHub Pages(Linux)は区別するため、ここではディレクトリ一覧との
    厳密な文字列一致で実在確認する。
    """
    if not dir_path.exists():
        return set()
    return {p.name for p in dir_path.iterdir()}


def check_media_file(filename, dir_names, dir_label, row_no, field_name, warnings):
    """ファイルの実在チェック（大文字小文字を区別する厳密一致）。"""
    if not filename:
        return
    if filename in dir_names:
        return
    lower_target = filename.lower()
    found_ci = next((name for name in dir_names if name.lower() == lower_target), None)
    if found_ci:
        warnings.append(
            f"行{row_no}: {field_name}「{filename}」は大文字小文字違いで"
            f"「{found_ci}」として存在します。GitHub Pagesは大文字小文字を区別するため"
            f"stories.xlsxかファイル名を修正してください。"
        )
    else:
        warnings.append(
            f"行{row_no}: {field_name}「{filename}」が{dir_label}/ に見つかりません。"
        )


def main():
    warnings = []

    if not XLSX_PATH.exists():
        print(f"エラー: {XLSX_PATH} が見つかりません。", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"エラー: シート「{SHEET_NAME}」が見つかりません。シート一覧: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)
    ws = wb[SHEET_NAME]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("エラー: シートが空です。", file=sys.stderr)
        sys.exit(1)

    header = [cell_str(h) for h in rows[0]]
    if header[: len(EXPECTED_HEADER)] != EXPECTED_HEADER:
        print(
            "警告: ヘッダー行が想定と異なります。\n"
            f"  想定: {EXPECTED_HEADER}\n"
            f"  実際: {header}",
            file=sys.stderr,
        )

    spots = []
    seen_ids = {}
    image_names = list_dir_names(IMAGES_DIR)
    audio_names = list_dir_names(AUDIO_DIR)

    for row_no, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None for v in row):
            continue  # 完全な空行はスキップ

        (
            id_raw, year_raw, theme_raw, title_raw, author_raw, place_raw,
            body_raw, lng_raw, lat_raw, photo1_raw, photo2_raw, photo3_raw,
            audio_raw, _remarks,
        ) = (list(row) + [None] * 14)[:14]

        if id_raw is None:
            warnings.append(f"行{row_no}: IDが空のためこの行をスキップします。")
            continue
        try:
            spot_id = int(id_raw)
        except (TypeError, ValueError):
            warnings.append(f"行{row_no}: IDが整数に変換できません（値: {id_raw!r}）。この行をスキップします。")
            continue

        if spot_id in seen_ids:
            warnings.append(f"行{row_no}: ID {spot_id} は行{seen_ids[spot_id]}と重複しています。")
        else:
            seen_ids[spot_id] = row_no

        year = parse_year(year_raw, row_no, warnings)

        try:
            lat = float(lat_raw)
            lng = float(lng_raw)
        except (TypeError, ValueError):
            warnings.append(f"行{row_no} (ID {spot_id}): 緯度または経度が数値ではありません。この行をスキップします。")
            continue

        if not (LAT_MIN <= lat <= LAT_MAX) or not (LNG_MIN <= lng <= LNG_MAX):
            warnings.append(
                f"行{row_no} (ID {spot_id}): 緯度経度が文京区周辺の想定範囲外です "
                f"(lat={lat}, lng={lng})。"
            )

        images = []
        for field_name, photo_raw in (
            ("写真1", photo1_raw), ("写真2", photo2_raw), ("写真3", photo3_raw),
        ):
            cell_value = cell_str(photo_raw)
            if not cell_value:
                continue
            # 1枚のセルにカンマ区切りで複数ファイル名を入れることも許容する
            # （写真欄が3つしかないため、3枚を超えるスポット向け）
            for name in [n.strip() for n in cell_value.split(",") if n.strip()]:
                check_media_file(name, image_names, "images", row_no, field_name, warnings)
                images.append(f"images/{name}")

        audio_name = cell_str(audio_raw)
        audio = None
        if audio_name:
            check_media_file(audio_name, audio_names, "audio", row_no, "音声", warnings)
            audio = f"audio/{audio_name}"

        spots.append({
            "id": spot_id,
            "year": year,
            "theme": cell_str(theme_raw),
            "title": cell_str(title_raw),
            "author": cell_str(author_raw),
            "place": cell_str(place_raw),
            "body": str(body_raw) if body_raw is not None else "",
            "lat": lat,
            "lng": lng,
            "images": images,
            "audio": audio,
        })

    spots.sort(key=lambda s: s["id"])

    OUTPUT_PATH.write_text(
        json.dumps(spots, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"{len(spots)}件のスポットを {OUTPUT_PATH.relative_to(ROOT)} に出力しました。")

    if warnings:
        print(f"\n警告 {len(warnings)}件:", file=sys.stderr)
        for w in warnings:
            print(f"  - {w}", file=sys.stderr)
    else:
        print("警告はありません。")


if __name__ == "__main__":
    main()
